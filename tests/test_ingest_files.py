import io

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from helpers import assert_order, zip_bytes

from careerforge.ingest.docx_reader import extract_docx
from careerforge.ingest.errors import (
    EmptyDocument,
    FileTooLarge,
    UnreadableDocument,
    UnsupportedFormat,
)
from careerforge.ingest.files import MAX_FILE_BYTES, read_document, sniff


def make_docx(build) -> bytes:
    doc = Document()
    build(doc)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def add_link(paragraph, url, text):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


class TestDocx:
    def test_paragraphs_tables_and_header_in_order(self):
        def build(d):
            d.sections[0].header.paragraphs[0].text = "PRIYA NAIR  priya@example.com"
            d.add_paragraph("EXPERIENCE")
            d.add_paragraph("Engineer at Acme")
            t = d.add_table(rows=2, cols=2)
            t.cell(0, 0).text, t.cell(0, 1).text = "Python", "Go"
            t.cell(1, 0).text, t.cell(1, 1).text = "SQL", "Redis"
            d.add_paragraph("EDUCATION")

        text, _ = extract_docx(make_docx(build))
        assert_order(text, ["PRIYA NAIR", "EXPERIENCE", "Engineer at Acme", "Python | Go", "SQL | Redis", "EDUCATION"])

    def test_hyperlinks_captured(self):
        def build(d):
            p = d.add_paragraph("Find me: ")
            add_link(p, "https://github.com/priya-demo", "GitHub")
            add_link(p, "mailto:a@b.co", "mail")

        text, links = extract_docx(make_docx(build))
        assert links == ["https://github.com/priya-demo"] and "GitHub" in text

    def test_merged_cells_do_not_repeat(self):
        def build(d):
            t = d.add_table(rows=1, cols=2)
            t.cell(0, 0).merge(t.cell(0, 1)).text = "Wide cell"

        text, _ = extract_docx(make_docx(build))
        assert text.count("Wide cell") == 1

    def test_text_box_content_is_found(self):
        def build(d):
            d.add_paragraph("Body paragraph")
            box = OxmlElement("w:txbxContent")
            p = OxmlElement("w:p")
            r = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = "SIDEBAR SKILLS Python"
            r.append(t)
            p.append(r)
            box.append(p)
            d.element.body.append(box)

        text, _ = extract_docx(make_docx(build))
        assert "SIDEBAR SKILLS Python" in text

    def test_macro_docm_style_archive_refused(self):
        data = zip_bytes({"word/document.xml": b"<x/>", "word/vbaProject.bin": b"\x00"})
        with pytest.raises(UnreadableDocument, match="Macro"):
            extract_docx(data)

    def test_zip_bomb_refused_before_parsing(self):
        data = zip_bytes({"word/document.xml": b"<x/>", "junk.bin": b"\x00" * 70_000_000})
        assert len(data) < 1_000_000  # compresses tiny: the dangerous kind
        with pytest.raises(UnreadableDocument, match="unusually large"):
            extract_docx(data)

    def test_zip_that_is_not_docx(self):
        with pytest.raises(UnreadableDocument, match="does not look like"):
            extract_docx(zip_bytes({"readme.txt": b"hi"}))

    def test_corrupt_zip(self):
        with pytest.raises(UnreadableDocument):
            extract_docx(b"PK\x03\x04 broken broken broken")


class TestReadDocument:
    def test_docx_via_dispatcher(self):
        data = make_docx(lambda d: d.add_paragraph("Priya Nair is a backend engineer with seven years of experience."))
        doc = read_document("cv.docx", data)
        assert doc.kind == "docx" and "backend engineer" in doc.text

    def test_text_utf8_and_bom(self):
        doc = read_document("cv.txt", "﻿Zoë Müller, senior engineer with ten years of Python and Go experience".encode())
        assert doc.kind == "txt" and doc.text.startswith("Zoë") and not doc.warnings

    def test_text_cp1252_is_decoded_with_a_warning(self):
        doc = read_document(
            "cv.txt", "Résumé of José, engineer with many years of Python and Go experience".encode("cp1252")
        )
        assert "José" in doc.text and doc.warnings

    def test_markdown_ok(self):
        assert (
            read_document("cv.md", b"# Priya Nair\n\nBackend engineer with seven years of experience in Python").kind
            == "txt"
        )

    @pytest.mark.parametrize("name,data,exc", [
        ("cv.pdf", b"", EmptyDocument),
        ("cv.txt", b"hi", EmptyDocument),
        ("cv.txt", b"x" * (MAX_FILE_BYTES + 1), FileTooLarge),
        ("cv.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"x" * 100, UnsupportedFormat),
        ("cv.pdf", b"This is plain text pretending to be a PDF " * 5, UnsupportedFormat),
        ("cv.docx", b"This is plain text pretending to be a docx " * 5, UnsupportedFormat),
        ("cv.txt", b"abc\x00\x01\x02binary" * 20, UnreadableDocument),
        ("cv.png", b"\x89PNG\r\n\x1a\n" + b"x" * 100, UnsupportedFormat),
        ("sheet.xlsx", zip_bytes({"x": b"y"}), UnsupportedFormat),
        ("cv.docm", zip_bytes({"word/document.xml": b"x"}), UnsupportedFormat),
    ])  # fmt: skip
    def test_bad_inputs_get_specific_errors(self, name, data, exc):
        with pytest.raises(exc):
            read_document(name, data)

    def test_profile_json_is_recognised_not_parsed_as_text(self):
        doc = read_document("me.json", b'{"kind": "careerforge.profile", "schema_version": 1, "profile": {}}')
        assert doc.kind == "profile_json" and doc.text == "" and doc.raw

    def test_other_json_is_rejected(self):
        with pytest.raises(UnsupportedFormat, match="not a CareerForge"):
            read_document("data.json", b'{"hello": "world"}')

    def test_zip_with_csvs_is_a_linkedin_export(self):
        data = zip_bytes({"Profile.csv": b"First Name\nA", "Positions.csv": b"Title\nB"})
        assert sniff("Basic_LinkedInDataExport.zip", data) == "linkedin_export"
        assert read_document("export.zip", data).kind == "linkedin_export"

    def test_filename_tricks_do_not_matter(self):
        assert sniff("../../etc/passwd.pdf", b"%PDF-1.7 x") == "pdf"
