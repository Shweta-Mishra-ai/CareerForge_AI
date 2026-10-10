"""CV rendering: options, every template, PDF/DOCX/TXT exports, gating and hostile content."""

import io
import re

import docx
import pytest
from bs4 import BeautifulSoup
from markupsafe import escape
from pydantic import ValidationError
from pypdf import PdfReader
from test_grounding import faithful_standard

from careerforge.models import (
    Certification,
    Contact,
    DateRange,
    Education,
    Experience,
    Profile,
    Project,
)
from careerforge.readiness import NotReady
from careerforge.render import cv
from careerforge.render.docx_export import render_cv_docx
from careerforge.render.options import TEMPLATES, RenderOptions, contrast_ratio
from careerforge.render.text_export import render_cv_txt
from careerforge.render.view import GROUP_MIN_SKILLS, build_blocks, skill_groups

ATS_SAFE = [k for k, t in TEMPLATES.items() if t.ats_safe]
ALL_TEMPLATES = list(TEMPLATES)


def ready(p: Profile) -> Profile:
    """The shared fixtures use example.com, which readiness (rightly) refuses to export."""
    return p.model_copy(update={"contact": p.contact.model_copy(update={"email": "priya.nair@gmail.com"})})


def experienced() -> Profile:
    return ready(faithful_standard())


def fresher() -> Profile:
    return Profile(
        name="Arjun Mehta",
        headline="Aspiring Full-Stack Developer",
        contact=Contact(email="arjun.mehta@gmail.com", location="Moradabad, India", github="github.com/arjun-demo"),
        skills=["Python", "JavaScript", "React", "Node.js", "MongoDB", "Git"],
        education=[
            Education(
                degree="B.Tech Information Technology",
                institution="Demo Institute of Technology",
                dates=DateRange(start="2021", end="2025"),
                grade="8.1 CGPA",
            )
        ],
        projects=[
            Project(
                name="Expense Tracker",
                tech=["React", "Node.js", "MongoDB"],
                url="https://github.com/arjun-demo/expense-tracker",
                bullets=["Built a full-stack web app to track expenses with charts."],
            )
        ],
    )


def maximal() -> Profile:
    jobs = [
        Experience(
            title=f"Engineer {i}",
            company=f"Company {i}",
            location="Remote",
            dates=DateRange(start=f"{2000 + i * 2}-01", end=f"{2001 + i * 2}-12"),
            bullets=[
                f"Delivered measurable result number {j} for the platform team, cutting cost by {j + 5}%."
                for j in range(8)
            ],
        )
        for i in range(8)
    ]
    return ready(faithful_standard()).model_copy(update={"experience": jobs})


PROFILES = {"experienced": experienced, "fresher": fresher}


def soup(html: str) -> BeautifulSoup:
    return BeautifulSoup(html, "html.parser")


# ── options ─────────────────────────────────────────────────────────────────


class TestOptions:
    @pytest.mark.parametrize(
        "bad",
        ["red", "#fff", "#12345", "#gggggg", "#1f4e79;}body{display:none", "url(http://evil)", '#1f4e79"><script>', ""],
    )
    def test_accent_cannot_carry_css_or_markup(self, bad):
        with pytest.raises(ValidationError):
            RenderOptions(accent=bad)

    def test_pale_accent_is_rejected_because_headings_would_be_unreadable(self):
        with pytest.raises(ValidationError, match="too light"):
            RenderOptions(accent="#ffff00")
        assert RenderOptions(accent="#1F4E79").accent == "#1f4e79"

    def test_contrast_ratio_matches_wcag_reference_values(self):
        assert contrast_ratio("#000000", "#ffffff") == pytest.approx(21.0)
        assert contrast_ratio("#ffffff", "#ffffff") == pytest.approx(1.0)
        assert contrast_ratio("#777777", "#ffffff") == pytest.approx(4.48, abs=0.01)  # the classic AA borderline

    @pytest.mark.parametrize(
        "kwargs", [{"template": "fancy"}, {"paper": "A3"}, {"font": "comic"}, {"density": "tight"}]
    )
    def test_unknown_choices_are_rejected(self, kwargs):
        with pytest.raises(ValidationError):
            RenderOptions(**kwargs)

    def test_order_rejects_duplicates_and_unknown_sections(self):
        with pytest.raises(ValidationError):
            RenderOptions(order=("skills", "skills"))
        with pytest.raises(ValidationError):
            RenderOptions(order=("hobbies",))

    @pytest.mark.parametrize("lang", ["en", "en-GB", "hi", "zh-Hans"])
    def test_good_language_codes(self, lang):
        assert RenderOptions(lang=lang).lang == lang

    @pytest.mark.parametrize("lang", ['en"><script>', "english", "", "e"])
    def test_bad_language_codes(self, lang):
        with pytest.raises(ValidationError):
            RenderOptions(lang=lang)

    def test_experienced_people_lead_with_experience_and_students_with_education(self):
        assert RenderOptions().section_order(experienced())[:2] == ["summary", "experience"]
        order = RenderOptions().section_order(fresher())
        assert order.index("education") < order.index("projects") < order.index("skills")

    def test_empty_and_hidden_sections_are_dropped(self):
        p = fresher()
        assert "experience" not in RenderOptions().section_order(p) and "summary" not in RenderOptions().section_order(
            p
        )
        assert "skills" not in RenderOptions(hidden=frozenset({"skills"})).section_order(p)

    def test_custom_order_is_respected_and_unlisted_sections_follow(self):
        order = RenderOptions(order=("skills", "education")).section_order(experienced())
        assert order[:2] == ["skills", "education"] and set(order) >= {"experience", "summary"}

    def test_options_are_immutable(self):
        with pytest.raises(ValidationError):
            RenderOptions().accent = "#000000"


# ── shared structure ────────────────────────────────────────────────────────


class TestBlocks:
    def test_blocks_follow_the_option_order_and_use_ats_standard_headings(self):
        titles = [b.title for b in build_blocks(experienced(), RenderOptions())]
        assert titles == ["Summary", "Work Experience", "Education", "Skills", "Certifications"]

    def test_entry_without_title_falls_back_to_company_and_does_not_repeat_it(self):
        p = Profile(
            name="A B", experience=[Experience(company="Acme", location="Pune", bullets=["Did the thing well"])]
        )
        (entry,) = build_blocks(p, RenderOptions())[0].entries
        assert entry.head == "Acme" and entry.sub == "Pune"

    def test_year_only_range_with_same_start_and_end_shows_one_year(self):
        assert DateRange(start="2018", end="2018").display() == "2018"
        assert DateRange(start="2018", end="2019").display() == "2018 – 2019"
        assert DateRange(start="2018-03", end="2018-09").display() == "Mar 2018 – Sep 2018"

    def test_skills_are_grouped_only_when_the_list_is_long_enough_and_categorised(self):
        few = ["Python", "Go", "Docker"]
        assert skill_groups(few, group=True) == [("", few)]
        many = ["Python", "Go", "Kafka", "FastAPI", "PostgreSQL", "Redis", "Docker", "Kubernetes", "AWS"]
        assert len(many) >= GROUP_MIN_SKILLS
        labels = [g for g, _ in skill_groups(many, group=True)]
        assert "Programming Languages" in labels and "Languages" not in labels  # not confusable with spoken languages
        assert skill_groups(many, group=False) == [("", many)]
        assert skill_groups(["Zumba"] * 0 + [f"Thing{i}" for i in range(12)], group=True)[0][0] == ""

    def test_no_skills_means_no_empty_group(self):
        assert skill_groups([], group=True) == []


# ── every template × every profile ─────────────────────────────────────────


@pytest.mark.parametrize("template", ALL_TEMPLATES)
@pytest.mark.parametrize("profile", PROFILES)
class TestEveryTemplate:
    def test_html_renders_with_all_facts_and_valid_structure(self, template, profile):
        p = PROFILES[profile]()
        html = cv.render_cv_html(p, RenderOptions(template=template))
        s = soup(html)
        assert html.startswith("<!DOCTYPE html>") and s.html["lang"] == "en"
        assert [h.get_text() for h in s.find_all("h1")] == [p.name]
        assert s.title.get_text() == f"{p.name} - CV"
        text = s.get_text(" ")
        for bullet in [b for e in p.experience for b in e.bullets] + [b for pr in p.projects for b in pr.bullets]:
            assert bullet in text
        assert not s.find(string=re.compile(r"\{\{|\{%"))  # no unrendered template syntax

    def test_pdf_has_selectable_text_and_correct_metadata(self, template, profile):
        p = PROFILES[profile]()
        out = cv.render_cv_pdf(p, RenderOptions(template=template))
        reader = PdfReader(io.BytesIO(out.data))
        assert len(reader.pages) == out.pages >= 1
        text = " ".join(page.extract_text() for page in reader.pages)
        assert p.name.casefold() in text.casefold() and p.contact.email in text  # executive prints the name in caps
        assert reader.metadata.title == f"{p.name} - CV" and reader.metadata.author == p.name


@pytest.mark.parametrize("template", ATS_SAFE)
@pytest.mark.parametrize("profile", PROFILES)
def test_ats_safe_templates_are_read_back_cleanly_from_the_actual_pdf(template, profile):
    out = cv.render_cv_pdf(PROFILES[profile](), RenderOptions(template=template))
    failed = [c.label for c in out.parse.checks if not c.passed]
    assert out.parse.score >= 95 and failed == [] and out.parse.findings == ()


def test_sidebar_template_is_honestly_flagged_not_called_safe():
    out = cv.render_cv_pdf(experienced(), RenderOptions(template="sidebar"))
    assert not TEMPLATES["sidebar"].ats_safe
    assert any("Side-by-side columns" in f for f in out.parse.findings)
    assert any("simple ATS parsers may mix the columns" in w for w in out.warnings)
    assert out.parse.score < cv.render_cv_pdf(experienced(), RenderOptions(template="ats-classic")).parse.score


def test_ruled_headings_are_not_mistaken_for_tables_by_the_parse_check():
    # regression: heading underlines are thin rectangles that pdfplumber used to read as a page-sized table
    out = cv.render_cv_pdf(experienced(), RenderOptions(template="ats-classic"))
    assert not any("Tables" in f for f in out.parse.findings)


def test_wide_letter_spacing_is_not_used_because_it_breaks_word_extraction():
    for name in TEMPLATES:
        css = (cv.TEMPLATE_DIR / "css" / f"{name}.css").read_text()
        for value in re.findall(r"letter-spacing:\s*([\d.]+)em", css):
            assert float(value) <= 0.05, f"{name}: letter-spacing {value}em splits words for some ATS parsers"


# ── layout options ──────────────────────────────────────────────────────────


class TestLayoutOptions:
    def test_paper_size_changes_the_pdf_page_box(self):
        def size(paper):
            box = (
                PdfReader(io.BytesIO(cv.render_cv_pdf(experienced(), RenderOptions(paper=paper)).data))
                .pages[0]
                .mediabox
            )
            return round(float(box.width)), round(float(box.height))

        assert size("A4") == (595, 842) and size("Letter") == (612, 792)

    def test_density_changes_length(self):
        p = maximal()
        compact = cv.render_cv_pdf(p, RenderOptions(density="compact")).pages
        relaxed = cv.render_cv_pdf(p, RenderOptions(density="relaxed")).pages
        assert compact < relaxed

    def test_long_cv_gets_a_page_count_warning_and_short_one_does_not(self):
        long = cv.render_cv_pdf(maximal(), RenderOptions(density="relaxed"))
        assert long.pages > 2 and any("pages" in w for w in long.warnings)
        assert cv.render_cv_pdf(experienced()).warnings == ()

    def test_accent_colour_reaches_the_css_and_font_stack_is_from_the_allowlist(self):
        html = cv.render_cv_html(experienced(), RenderOptions(template="ats-modern", accent="#8b0000", font="serif"))
        assert "--accent: #8b0000;" in html and "'Liberation Serif'" in html

    def test_hidden_section_and_custom_order(self):
        html = cv.render_cv_html(experienced(), RenderOptions(hidden=frozenset({"certifications"}), order=("skills",)))
        headings = [h.get_text() for h in soup(html).find_all("h2")]
        assert "Certifications" not in headings and headings[0] == "Skills"

    def test_sidebar_puts_skills_in_the_aside_and_experience_in_main(self):
        s = soup(cv.render_cv_html(experienced(), RenderOptions(template="sidebar")))
        assert "Skills" in [h.get_text() for h in s.aside.find_all("h2")]
        assert "Work Experience" in [h.get_text() for h in s.main.find_all("h2")]

    def test_sidebar_contact_is_in_the_aside_exactly_once(self):
        html = cv.render_cv_html(experienced(), RenderOptions(template="sidebar"))
        assert html.count("mailto:priya.nair@gmail.com") == 1

    def test_contact_links_are_real_links(self):
        s = soup(cv.render_cv_html(experienced()))
        hrefs = {a["href"] for a in s.select(".contact a")}
        assert "mailto:priya.nair@gmail.com" in hrefs and "tel:+919876543210" in hrefs
        assert "https://linkedin.com/in/priya-nair-demo" in hrefs

    def test_project_links_show_a_short_url(self):
        s = soup(cv.render_cv_html(fresher()))
        link = s.select_one(".sec-projects a")
        assert link["href"] == "https://github.com/arjun-demo/expense-tracker"
        assert link.get_text() == "github.com/arjun-demo/expense-tracker"


# ── readiness gate ──────────────────────────────────────────────────────────

EXPORTERS = {
    "html": cv.render_cv_html,
    "pdf": cv.render_cv_pdf,
    "docx": render_cv_docx,
    "txt": render_cv_txt,
}


class TestGate:
    @pytest.mark.parametrize("fmt", EXPORTERS)
    def test_sample_email_blocks_every_export(self, fmt):
        with pytest.raises(NotReady, match="sample email"):
            EXPORTERS[fmt](faithful_standard())

    @pytest.mark.parametrize("fmt", EXPORTERS)
    def test_nothing_to_show_blocks_every_export(self, fmt):
        with pytest.raises(NotReady):
            EXPORTERS[fmt](Profile(name="Real Person", contact=Contact(email="real.person@gmail.com")))

    @pytest.mark.parametrize("fmt", EXPORTERS)
    def test_leftover_placeholder_blocks_every_export(self, fmt):
        p = experienced()
        bad = p.experience[0].model_copy(update={"bullets": ["Achieved [mention specific achievement] for the team"]})
        with pytest.raises(NotReady, match="placeholder"):
            EXPORTERS[fmt](p.model_copy(update={"experience": [bad, *p.experience[1:]]}))

    @pytest.mark.parametrize("fmt", EXPORTERS)
    def test_unconfirmed_ai_suggestions_never_reach_any_export(self, fmt):
        p = experienced()
        ghost = Experience(title="Wizard", company="FakeCo", origin="suggested", bullets=["Conjured FakeCo revenue"])
        cert = Certification(name="Fake Cert", origin="suggested")
        p = p.model_copy(update={"experience": [ghost, *p.experience], "certifications": [cert, *p.certifications]})
        out = EXPORTERS[fmt](p)
        if isinstance(out, cv.PdfExport):
            blob = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(out.data)).pages)
        elif isinstance(out, bytes):
            blob = "\n".join(par.text for par in docx.Document(io.BytesIO(out)).paragraphs)
        else:
            blob = out
        assert "FakeCo" not in blob and "Fake Cert" not in blob and "Wizard" not in blob
        assert "Finlytics" in blob  # the real content is still there (guards against an over-eager filter)

    def test_confirmed_suggestions_are_exported(self):
        p = experienced()
        accepted = Experience(
            title="Intern", company="Acme", origin="suggested", confirmed=True, bullets=["Wrote docs"]
        )
        assert "Acme" in render_cv_txt(p.model_copy(update={"experience": [*p.experience, accepted]}))

    def test_preview_works_for_a_draft_and_says_it_is_a_draft(self):
        draft = faithful_standard()  # sample email: not ready
        html = cv.preview_html(draft)
        assert "Draft preview" in html and "Priya Nair" in html

    def test_preview_of_a_ready_profile_has_no_banner(self):
        assert "Draft preview" not in cv.preview_html(experienced())

    def test_preview_also_hides_unconfirmed_suggestions(self):
        ghost = Experience(title="Wizard", company="FakeCo", origin="suggested", bullets=["Conjured revenue"])
        assert "FakeCo" not in cv.preview_html(experienced().model_copy(update={"experience": [ghost]}))

    def test_input_profile_is_not_modified(self):
        p = experienced()
        before = p.model_dump()
        for fn in EXPORTERS.values():
            fn(p)
        assert p.model_dump() == before


# ── hostile content ─────────────────────────────────────────────────────────

# Passes readiness (no complete tag) but would still corrupt or hijack the page if it were not escaped.
SNEAKY = [
    "5 < 6 && 7 > 3 \"quoted\" 'single' & friends",
    "a </style b",
    "<img src=x onerror=window.__pwned=1",
    '" onmouseover="window.__pwned=1',
    "x </title b </head c",
    "&lt;script&gt; &amp; &#x3c;",
]
# Complete markup: readiness refuses it for export, but `preview_html` must still neutralise it.
FULL = [
    "<script>window.__pwned=1</script>",
    '<img src=x onerror="window.__pwned=1">',
    '"><svg/onload=window.__pwned=1>',
    "<iframe src=javascript:alert(1)></iframe>",
    "{{7*7}} {% for x in y %}{% endfor %} ${7*7} #{7*7}",
]


def hostile(value: str) -> Profile:
    return Profile(
        name=f"Eve {value}"[:120],
        headline=value[:200],
        summary=value,
        contact=Contact(email="eve@gmail.com", location=value[:100]),
        skills=["Python", value[:60]],
        languages=[value[:60]],
        experience=[
            Experience(title=value[:150], company=value[:150], location=value[:100], bullets=[value, value + " second"])
        ],
        education=[Education(degree=value[:200], institution=value[:200], grade=value[:60], details=[value])],
        projects=[Project(name=value[:150], description=value, tech=[value[:60]], bullets=[value])],
        certifications=[Certification(name=value[:200], issuer=value[:150])],
    )


def assert_inert(html: str, payload: str):
    s = soup(html)
    assert s.find_all("script") == []  # CVs contain no script at all
    assert not s.find(["img", "iframe", "svg", "object", "embed", "link", "base", "form"])
    assert not [t.name for t in s.find_all(True) if any(a.lower().startswith("on") for a in t.attrs)]
    assert len(s.find_all("style")) == 1  # a </style> in the content did not close the sheet or open another
    assert str(escape(payload)) in html  # exactly the escaped form is in the markup
    assert payload in s.find("section", class_="sec-summary").get_text()  # and a browser shows what was typed


@pytest.mark.parametrize("template", ALL_TEMPLATES)
@pytest.mark.parametrize("payload", SNEAKY)
def test_sneaky_text_is_escaped_in_the_final_html(template, payload):
    assert_inert(cv.render_cv_html(hostile(payload), RenderOptions(template=template)), payload)


@pytest.mark.parametrize("template", ALL_TEMPLATES)
@pytest.mark.parametrize("payload", FULL)
def test_full_markup_is_escaped_even_in_the_ungated_preview(template, payload):
    assert_inert(cv.preview_html(hostile(payload), RenderOptions(template=template)), payload)


@pytest.mark.parametrize("payload", FULL)
def test_gate_refuses_to_export_full_markup(payload):
    with pytest.raises(NotReady):
        cv.render_cv_html(hostile(payload))


@pytest.mark.parametrize("bad", ["javascript:alert(1)", "data:text/html,<b>x</b>", "vbscript:x", "file:///etc/passwd"])
def test_links_that_are_not_http_cannot_exist_on_a_profile(bad):
    with pytest.raises(ValueError):
        Contact(website=bad)
    with pytest.raises(ValueError):
        Project(name="x", url=bad)


def test_pdf_generation_never_connects_to_anything():
    import socket
    import threading

    from weasyprint import HTML

    server = socket.socket()
    server.bind(("127.0.0.1", 0))
    server.listen(5)
    server.settimeout(0.5)
    port, hits = server.getsockname()[1], []

    def accept():
        try:
            conn, _ = server.accept()
            hits.append(1)
            conn.close()
        except TimeoutError:
            pass

    t = threading.Thread(target=accept)
    t.start()
    html = f'<link rel="stylesheet" href="http://127.0.0.1:{port}/a.css"><img src="http://127.0.0.1:{port}/b.png"><p>ok</p>'
    assert HTML(string=html, url_fetcher=cv.NO_FETCH).write_pdf()  # renders, with the resources skipped
    t.join()
    server.close()
    assert hits == []
    # positive control: the same document with WeasyPrint's default fetcher does reach the socket
    server = socket.socket()
    server.bind(("127.0.0.1", 0))
    server.listen(5)
    server.settimeout(2)
    port, hits = server.getsockname()[1], []
    t = threading.Thread(target=accept)
    t.start()
    HTML(string=f'<img src="http://127.0.0.1:{port}/b.png">').write_pdf()
    t.join()
    server.close()
    assert hits == [1]


@pytest.mark.parametrize("name", ["李小龍", "محمد علي", "Zoë Ångström-Łukasiewicz", "A" * 120])
def test_unusual_names_render_in_every_format(name):
    p = experienced().model_copy(update={"name": name})
    assert name in cv.render_cv_html(p)
    assert name.upper() in render_cv_txt(p)
    assert render_cv_pdf_ok(p)
    assert docx.Document(io.BytesIO(render_cv_docx(p))).paragraphs[0].text == name


def render_cv_pdf_ok(p: Profile) -> bool:
    return len(cv.render_cv_pdf(p).data) > 1000


def test_maximum_size_profile_renders_in_every_format():
    p = maximal()
    assert cv.render_cv_pdf(p).pages >= 2
    assert len(render_cv_docx(p)) > 5000 and render_cv_txt(p).count("\n- ") == 64


# ── TXT ─────────────────────────────────────────────────────────────────────


class TestTxt:
    def test_layout_and_content(self):
        t = render_cv_txt(experienced())
        lines = t.splitlines()
        assert lines[0] == "PRIYA NAIR" and lines[1] == "Senior Backend Engineer"
        assert "priya.nair@gmail.com | +91 98765 43210 | Bengaluru, India" in lines[2]
        assert "WORK EXPERIENCE\n---------------" in t
        assert "Senior Backend Engineer  (Mar 2021 – Present)" in t and "Finlytics Pvt Ltd" in t
        assert "- Cut p95 API latency from 480ms to 190ms by introducing Redis caching." in t
        assert "Programming Languages: Python, Go" in t
        assert t.endswith("\n") and not t.endswith("\n\n") and "\n\n\n" not in t

    def test_section_order_option_applies(self):
        t = render_cv_txt(experienced(), RenderOptions(order=("skills",)))
        assert t.index("SKILLS") < t.index("WORK EXPERIENCE")

    def test_every_bullet_appears_exactly_once(self):
        p = experienced()
        t = render_cv_txt(p)
        for b in (b for e in p.experience for b in e.bullets):
            assert t.count(b) == 1


# ── DOCX ────────────────────────────────────────────────────────────────────


class TestDocx:
    def doc(self, p=None, **kw):
        return docx.Document(io.BytesIO(render_cv_docx(p or experienced(), RenderOptions(**kw))))

    def test_uses_real_word_structure_and_no_tables_or_images(self):
        d = self.doc()
        assert len(d.tables) == 0 and len(d.inline_shapes) == 0
        styles = [p.style.name for p in d.paragraphs]
        assert styles.count("Heading 1") == 5  # one per section, so Word's navigation pane works
        assert styles.count("List Bullet") == 5
        assert d.paragraphs[0].style.name == "CV Name"

    def test_text_is_in_reading_order(self):
        text = "\n".join(p.text for p in self.doc().paragraphs)
        positions = [
            text.index(x) for x in ("Priya Nair", "SUMMARY", "WORK EXPERIENCE", "EDUCATION", "SKILLS", "CERTIFICATIONS")
        ]
        assert positions == sorted(positions)

    def test_dates_use_a_right_aligned_tab_stop_not_a_table(self):
        d = self.doc()
        row = next(p for p in d.paragraphs if p.text.startswith("Senior Backend Engineer\t"))
        assert row.text.endswith("Mar 2021 – Present")
        stops = row.style.paragraph_format.tab_stops
        assert len(stops) == 1 and stops[0].alignment == 2  # WD_TAB_ALIGNMENT.RIGHT

    def test_document_properties_identify_the_candidate_not_the_library(self):
        props = self.doc().core_properties
        assert (
            props.author == "Priya Nair" and props.title == "Priya Nair - CV" and props.last_modified_by == "Priya Nair"
        )
        assert props.language == "en" and "python-docx" not in (props.author + props.comments + props.subject)

    @pytest.mark.parametrize("paper,width", [("A4", 210), ("Letter", 215.9)])
    def test_paper_size(self, paper, width):
        assert self.doc(paper=paper).sections[0].page_width.mm == pytest.approx(width, abs=0.1)

    def test_font_follows_the_option(self):
        assert self.doc(font="serif").styles["Normal"].font.name == "Times New Roman"
        assert self.doc(font="humanist").styles["Normal"].font.name == "Calibri"

    def test_classic_template_has_black_headings_and_modern_uses_the_accent(self):
        assert self.doc(template="ats-classic").styles["Heading 1"].font.color.rgb == docx.shared.RGBColor(0, 0, 0)
        assert str(self.doc(template="ats-modern", accent="#8b0000").styles["Heading 1"].font.color.rgb) == "8B0000"

    def test_hostile_text_is_stored_as_text(self):
        for payload in SNEAKY:
            d = self.doc(hostile(payload))
            assert payload in "\n".join(p.text for p in d.paragraphs)

    def test_output_is_a_valid_zip_with_expected_parts(self):
        import zipfile

        with zipfile.ZipFile(io.BytesIO(render_cv_docx(experienced()))) as z:
            assert z.testzip() is None and "word/document.xml" in z.namelist()


def test_template_typos_fail_loudly_instead_of_printing_nothing():
    from jinja2 import UndefinedError

    with pytest.raises(UndefinedError):
        cv._env().from_string("{{ no_such_variable.name }}").render()


def test_fresher_without_jobs_is_not_penalised_for_a_missing_experience_heading():
    out = cv.render_cv_pdf(fresher(), RenderOptions())
    assert not any("experience" in f for f in out.parse.findings)
