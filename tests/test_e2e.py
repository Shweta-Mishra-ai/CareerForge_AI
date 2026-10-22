"""Real browser, real Streamlit server, real downloads. Marked e2e; needs Chromium."""

import io
import json
import re
import zipfile

import pdfplumber
import pytest
from e2e_fixtures import CV, JD, PROFILE, Session
from playwright.sync_api import expect
from test_writing import GOOD_PARAS

from careerforge.ats.checks import check_pdf
from careerforge.models import Profile

pytestmark = pytest.mark.e2e

# ── journeys without AI ─────────────────────────────────────────────────────


def test_old_cv_to_a_submittable_cv_and_portfolio(session):
    """Upload an old CV, fix a detail, check it against a job, download PDF/DOCX/TXT and a portfolio."""
    s, page = session, session.page
    s.load_cv()
    expect(page.get_by_label("Full name")).to_have_value("Priya Nair")
    expect(page.get_by_label("Email")).to_have_value("priya.nair@gmail.com")
    expect(page.get_by_text("Ready to export").first).to_be_visible()

    # edit the headline and save
    page.get_by_label("Headline (e.g. Senior Backend Engineer)").fill("Staff Backend Engineer")
    s.click("Save basics")
    expect(page.get_by_label("Headline (e.g. Senior Backend Engineer)")).to_have_value("Staff Backend Engineer")

    # ATS check against a posting
    s.go("ATS check")
    s.paste_job()
    expect(page.get_by_text("ATS match score")).to_be_visible()
    expect(page.get_by_text("REST API").first).to_be_visible()  # a requirement the CV does not show

    # CV export
    s.go("CV")
    s.click("Build my files")
    expect(page.get_by_text("A simple ATS text extractor reads this PDF cleanly").first).to_be_visible(timeout=60000)
    pdf = s.download("PDF").path()
    data = open(pdf, "rb").read()
    with pdfplumber.open(io.BytesIO(data)) as doc:
        text = "\n".join(p.extract_text() or "" for p in doc.pages)
    assert "Priya Nair" in text and "Staff Backend Engineer" in text and "Finlytics" in text
    assert check_pdf(data, PROFILE).score >= 95

    import docx

    document = docx.Document(s.download("Word (.docx)").path())
    assert document.paragraphs[0].text == "Priya Nair" and not document.tables
    assert "WORK EXPERIENCE" in open(s.download("Plain text").path(), encoding="utf-8").read()

    # portfolio
    s.go("Portfolio")
    expect(page.get_by_role("button", name="Download site (.zip with deploy guide)")).to_be_visible(timeout=30000)
    z = zipfile.ZipFile(s.download("Download site (.zip with deploy guide)").path())
    html = z.read("index.html").decode()
    assert "Priya Nair" in html and "98765" not in html and "<script src" not in html
    s.assert_clean()


def test_portfolio_v2_options_reach_the_downloaded_site(session):
    """Automatic theme, the CV download button, a custom domain and no-index: all land in the real zip."""
    s, page = session, session.page
    s.load_cv()
    s.go("Portfolio")
    s.choose("Automatic")
    domain = page.get_by_role("textbox", name="Own domain (optional)")
    domain.fill("jane.dev")
    domain.blur()
    page.get_by_text("Add a CV download button").click()
    page.get_by_text("Hide from search engines").click()
    expect(page.get_by_role("button", name="Download site (.zip with deploy guide)")).to_be_visible(timeout=60000)
    page.wait_for_timeout(1500)
    z = zipfile.ZipFile(s.download("Download site (.zip with deploy guide)").path())
    names = set(z.namelist())
    assert {"index.html", "404.html", "robots.txt", "CNAME", "Priya_Nair_CV.pdf", "README.md", ".nojekyll"} <= names
    assert z.read("CNAME") == b"jane.dev\n" and z.read("robots.txt") == b"User-agent: *\nDisallow: /\n"
    html = z.read("index.html").decode()
    assert "prefers-color-scheme: dark" in html and 'name="robots" content="noindex"' in html
    assert 'href="Priya_Nair_CV.pdf" download' in html and z.read("Priya_Nair_CV.pdf")[:5] == b"%PDF-"
    s.assert_clean()


def test_spotlight_with_a_photo_and_the_badge_reaches_the_downloaded_site(session, tmp_path):
    """Upload a photo in the real page; the zip's index.html carries it inline, re-saved, with the badge."""
    import base64
    import io

    from PIL import Image

    s, page = session, session.page
    s.load_cv()
    s.go("Portfolio")
    page.get_by_test_id("stRadio").filter(has_text="Layout").get_by_text("Spotlight", exact=True).click()
    page.get_by_text("Open to opportunities").click()
    exif = Image.Exif()
    exif[0x010F] = "PhoneMaker"
    photo = tmp_path / "me.jpg"
    Image.new("RGB", (900, 1200), (180, 140, 120)).save(photo, "JPEG", exif=exif.tobytes())
    page.get_by_test_id("stFileUploader").filter(has_text="Your photo").locator("input[type=file]").set_input_files(
        str(photo)
    )
    expect(page.frame_locator("iframe").first.locator("img.avatar")).to_be_visible(timeout=60000)
    page.wait_for_timeout(1500)
    z = zipfile.ZipFile(s.download("Download site (.zip with deploy guide)").path())
    html = z.read("index.html").decode()
    assert 'class="spot"' in html and "Open to opportunities" in html
    src = html.split('src="data:image/jpeg;base64,', 1)[1].split('"', 1)[0]
    data = base64.b64decode(src)
    assert b"PhoneMaker" not in data and Image.open(io.BytesIO(data)).size == (480, 480)
    s.assert_clean()


def test_a_file_that_is_not_a_photo_is_explained(session, tmp_path):
    s, page = session, session.page
    s.load_cv()
    s.go("Portfolio")
    fake = tmp_path / "me.png"
    fake.write_bytes(b"not really a picture")
    page.get_by_test_id("stFileUploader").filter(has_text="Your photo").locator("input[type=file]").set_input_files(
        str(fake)
    )
    expect(page.get_by_text("could not be read as a photo")).to_be_visible(timeout=30000)
    s.assert_clean()


def test_a_job_posting_can_be_uploaded_as_pdf_word_or_text(session, tmp_path):
    """The old CareerBoost app took the posting as an uploaded PDF; so must this one."""
    import docx
    from weasyprint import HTML

    s, page = session, session.page
    s.load_cv()
    s.go("ATS check")
    posting = "Senior Backend Engineer\nRequirements: Python, Kafka, Redis, Terraform and Kubernetes experience."
    txt = tmp_path / "posting.txt"
    txt.write_text(posting, encoding="utf-8")
    word = tmp_path / "posting.docx"
    d = docx.Document()
    d.add_paragraph(posting.replace("Python", "Rust"))
    d.save(word)
    pdf = tmp_path / "posting.pdf"
    HTML(string=f"<p>{posting.replace('Python', 'Elixir')}</p>").write_pdf(pdf)
    page.get_by_text("Or upload the posting (PDF, Word or text file)").click()
    for path, term in ((txt, "Python"), (word, "Rust"), (pdf, "Elixir")):
        page.locator("input[type=file]").first.set_input_files(str(path))
        expect(page.get_by_role("textbox", name="Job posting (paste the full text)")).to_have_value(
            re.compile(term), timeout=30000
        )
    expect(page.get_by_text("ATS match score")).to_be_visible(timeout=30000)
    # a scanned/empty or wrong kind of file is a sentence, not a crash
    junk = tmp_path / "junk.txt"
    junk.write_text("hi", encoding="utf-8")
    page.locator("input[type=file]").first.set_input_files(str(junk))
    expect(page.get_by_text("not enough text").first).to_be_visible(timeout=30000)
    s.assert_clean()


def test_a_saved_search_refills_the_boxes_in_the_browser(session):
    """Boxes must show what the app holds: the same 'looks blank but the app has the value' trap as the upload."""
    s, page = session, session.page
    s.load_cv()
    s.go("Jobs & tracker")
    query = page.get_by_role("textbox", name="Search", exact=True)
    where = page.get_by_role("textbox", name="Where (optional)")
    query.fill("data engineer")
    query.blur()
    where.fill("Pune")
    where.blur()
    page.wait_for_timeout(500)
    s.click("Save this search")
    expect(page.get_by_text("Saved searches")).to_be_visible(timeout=30000)
    query.fill("something else")
    query.blur()
    where.fill("")
    where.blur()
    page.wait_for_timeout(500)
    page.get_by_role("button", name="Load", exact=True).click()
    expect(query).to_have_value("data engineer", timeout=30000)
    expect(where).to_have_value("Pune")
    s.assert_clean()


def test_each_portfolio_layout_downloads_as_a_working_site(session, browser):
    """Pick a layout in the real UI, download the zip, open the downloaded index.html in a browser."""
    import tempfile

    s, page = session, session.page
    s.load_cv()
    s.go("Portfolio")
    for label, marker in (("Hero", "hero"), ("Editorial", "masthead"), ("Profile", "profile")):
        page.get_by_test_id("stRadio").filter(has_text="Layout").get_by_text(label, exact=True).click()
        expect(page.get_by_role("button", name="Download site (.zip with deploy guide)")).to_be_visible(timeout=60000)
        page.wait_for_timeout(1500)
        z = zipfile.ZipFile(s.download("Download site (.zip with deploy guide)").path())
        html = z.read("index.html").decode()
        assert f'class="{marker}"' in html, label
        with tempfile.TemporaryDirectory() as tmp:
            z.extractall(tmp)
            site = browser.new_page()
            site.goto(f"file://{tmp}/index.html")
            expect(site.get_by_role("heading", name="Priya Nair", level=1)).to_be_visible()
            assert site.evaluate("document.documentElement.scrollWidth") <= site.evaluate("innerWidth")
            site.close()
    s.assert_clean()


def test_save_load_round_trip_restores_everything(session, browser, server, tmp_path):
    s, page = session, session.page
    s.load_cv()
    s.go("Jobs & tracker")
    s.choose("Application tracker")
    page.get_by_label("Company", exact=True).fill("Acme")
    page.get_by_label("Role", exact=True).fill("Backend Engineer")
    s.click("Add application")
    expect(
        page.locator("td[role=gridcell]", has_text="Acme").first
    ).to_be_attached()  # the grid is a canvas; cells are in its a11y table
    download = s.download("Save my profile")
    saved = str(tmp_path / download.suggested_filename)
    download.save_as(saved)
    saved_text = open(saved, encoding="utf-8").read()
    envelope = json.loads(saved_text)
    assert envelope["kind"] == "careerforge.profile" and envelope["applications"][0]["company"] == "Acme"
    assert Profile.from_json(saved_text).name == "Priya Nair"

    fresh = Session(browser, server)  # a brand new session knows nothing
    try:
        expect(fresh.page.get_by_text("No profile loaded yet")).to_be_visible()
        fresh.page.set_input_files("input[type=file]", saved)
        fresh.click("Read this file")
        fresh.wait_title("Review & edit")
        expect(fresh.page.get_by_label("Full name")).to_have_value("Priya Nair")
        fresh.go("Jobs & tracker")
        fresh.choose("Application tracker")
        expect(fresh.page.locator("td[role=gridcell]", has_text="Acme").first).to_be_attached()
        fresh.assert_clean()
    finally:
        fresh.close()


def test_validation_errors_are_shown_and_nothing_is_lost(session, tmp_path):
    s, page = session, session.page
    s.load_cv()
    page.get_by_label("Email").fill("not-an-email")
    s.click("Save basics")
    expect(page.get_by_text("Not saved").first).to_be_visible()
    expect(page.get_by_label("Email")).to_have_value("not-an-email")  # what you typed stays, so you can fix it
    saved = s.download("Save my profile")  # ...but the profile itself still holds the last good value
    assert (
        json.loads(open(saved.path(), encoding="utf-8").read())["profile"]["contact"]["email"] == "priya.nair@gmail.com"
    )


def test_not_ready_cv_cannot_be_downloaded(session):
    s, page = session, session.page
    s.load_cv()
    page.get_by_label("Full name").fill("")
    s.click("Save basics")
    s.go("CV")
    expect(page.get_by_text("must be fixed before you can export")).to_be_visible()
    expect(page.get_by_role("button", name="Build my files")).to_have_count(0)
    frame = page.frame_locator("iframe").first
    expect(frame.get_by_text("Draft preview")).to_be_visible()


def test_pages_are_reachable_by_keyboard_alone(session):
    s, page = session, session.page
    s.load_cv()
    page.get_by_role("link", name="ATS check").first.focus()
    page.keyboard.press("Enter")
    s.wait_title("ATS check")
    page.get_by_role("textbox", name="Job posting (paste the full text)").focus()
    page.keyboard.type(JD)
    page.keyboard.press("Tab")
    page.wait_for_timeout(600)
    expect(page.get_by_text("ATS match score")).to_be_visible()


# ── journeys with (fake) AI ─────────────────────────────────────────────────


def test_ai_journey_read_tailor_letter_interview(ai_session, ai):
    s, page = ai_session, ai_session.page
    expect(page.get_by_text("AI: ready (30 shared actions left)")).to_be_visible()
    page.set_input_files("input[type=file]", str(CV))
    s.click("Read this file")
    s.wait_title("Review & edit")
    expect(page.get_by_label("Full name")).to_have_value("Priya Nair")
    assert any("data-extraction engine" in r["system"] for r in ai.requests)
    assert all(r["key"] == "fake-e2e-key-1234567890" for r in ai.requests)  # the configured key is what is sent

    s.go("Tailor to a job")
    s.paste_job()
    s.click("Suggest improvements")
    expect(page.get_by_text("Applied automatically")).to_be_visible(timeout=30000)
    expect(page.get_by_text("Needs your decision")).to_be_visible()
    page.get_by_text("This is true for me").first.click()
    s.click("Use this tailored version")
    expect(page.get_by_text("Your main profile is unchanged")).to_be_visible()

    s.go("CV")
    s.choose("Tailored version")
    preview = page.frame_locator("iframe").first.locator("body")
    expect(preview).to_contain_text("Reduced p95 API latency")
    expect(preview).to_contain_text("REST API endpoints")

    s.go("Letters & interviews")
    page.get_by_label("Company", exact=True).fill("Northwind Payments")
    s.click("Create my cover letter")
    expect(page.get_by_text("Checked: every number, name and skill")).to_be_visible(timeout=30000)
    box = page.get_by_label(re.compile("Letter body"))
    expect(box).to_have_value(re.compile("^" + re.escape(GOOD_PARAS[0][:40])))
    with pdfplumber.open(s.download("PDF").path()) as letter_pdf:
        letter_text = " ".join(pg.extract_text() or "" for pg in letter_pdf.pages)
        assert len(letter_pdf.pages) == 1 and "Dear Hiring Manager," in letter_text and "Priya Nair" in letter_text
    s.choose("Interview prep")
    s.click("Prepare my interview questions")
    expect(page.get_by_text("Technical: How did you reduce API latency at Finlytics?")).to_be_visible(timeout=30000)

    s.go("Settings")
    expect(page.get_by_text(re.compile(r"\d+ AI actions left in this session"))).to_be_visible()
    s.assert_clean()
    assert page.content().count("fake-e2e-key") == 0  # the key never reaches the browser


def test_provider_outage_is_explained_and_the_basic_reader_rescues_the_upload(ai_session, ai):
    s, page = ai_session, ai_session.page
    ai.fail = True
    page.set_input_files("input[type=file]", str(CV))
    s.click("Read this file")
    expect(page.get_by_text("could not be reached").first).to_be_visible(timeout=120000)
    s.click("Read it with the basic reader instead (no AI)")
    s.wait_title("Review & edit")
    expect(page.get_by_label("Full name")).to_have_value("Priya Nair")
    assert "fake-e2e-key" not in s.server.log_text()
    s.assert_clean()
