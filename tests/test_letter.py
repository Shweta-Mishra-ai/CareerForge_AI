"""Cover letter v2: business-letter format, evidence pairing, lint, the AI-free letter and rendering."""

import io
import re
from dataclasses import replace
from datetime import date

import docx
import pytest
from pypdf import PdfReader
from test_render_cv import experienced, fresher
from test_writing import GOOD_PARAS, JD, Scripted, client

from careerforge.ai.grounding import profile_text
from careerforge.ats.report import analyze
from careerforge.models import Experience, Profile, Project
from careerforge.render.letter import (
    LetterNotReady,
    render_letter_docx,
    render_letter_html,
    render_letter_pdf,
    render_letter_txt,
    with_body,
)
from careerforge.render.options import RenderOptions
from careerforge.writing.basic_letter import build_basic_letter
from careerforge.writing.claims import unsupported_claims
from careerforge.writing.cover_letter import SYSTEM, write_cover_letter
from careerforge.writing.evidence import Evidence, evidence_pairs
from careerforge.writing.letter_format import LetterDetails, assemble_letter, closing_for, format_date
from careerforge.writing.letter_lint import CLICHES, lint_letter

TODAY = date(2026, 10, 10)
P = experienced()
FULL = LetterDetails(
    company="Northwind Payments",
    recipient_name="Ms. Rao",
    recipient_title="Head of Engineering",
    company_city="Pune",
    job_ref="ENG-2041",
    source="LinkedIn",
    availability="I can start on 1 November",
    hook="Northwind's focus on reliable settlement is the kind of problem I enjoy.",
)


def checks(letter, **kw):
    return {c.key: c for c in lint_letter(letter, P, kw.pop("job", JD), **kw)}


class TestFormat:
    @pytest.mark.parametrize(
        "style,text", [("us", "October 10, 2026"), ("uk", "10 October 2026"), ("iso", "2026-10-10")]
    )
    def test_dates_follow_the_readers_country(self, style, text):
        assert format_date(TODAY, style) == text

    def test_single_digit_days_have_no_leading_zero(self):
        assert (
            format_date(date(2026, 3, 5), "us") == "March 5, 2026"
            and format_date(date(2026, 3, 5), "uk") == "5 March 2026"
        )

    @pytest.mark.parametrize(
        "style,tone,name,expected",
        [
            ("us", "standard", "Ms. Rao", "Sincerely,"),
            ("us", "standard", "", "Sincerely,"),
            ("uk", "standard", "Ms. Rao", "Yours sincerely,"),
            ("uk", "standard", "", "Yours faithfully,"),
            ("uk", "warm", "Ms. Rao", "Kind regards,"),
        ],
    )
    def test_closing_matches_style_and_whether_the_recipient_is_named(self, style, tone, name, expected):
        assert closing_for(LetterDetails(date_style=style, tone=tone, recipient_name=name)) == expected

    def test_full_letter_has_every_standard_part_in_order(self):
        letter = assemble_letter(P, ["Para one.", "Para two."], FULL, TODAY)
        text = letter.text()
        order = ["Priya Nair", "October 10, 2026", "Ms. Rao", "Head of Engineering", "Northwind Payments", "Pune",
                 "Re: Application", "Dear Ms. Rao,", "Para one.", "Para two.", "Sincerely,", "Enclosure: Curriculum Vitae"]  # fmt: skip
        positions = [text.index(x) for x in order]
        assert positions == sorted(positions)

    def test_subject_line_carries_role_and_reference_only_when_known(self):
        assert (
            assemble_letter(P, ["x"], replace(FULL, role="Backend Engineer"), TODAY).subject
            == "Re: Application for Backend Engineer (Ref: ENG-2041)"
        )
        assert (
            assemble_letter(P, ["x"], LetterDetails(role="Backend Engineer"), TODAY).subject
            == "Re: Application for Backend Engineer"
        )
        assert assemble_letter(P, ["x"], LetterDetails(job_ref="A1"), TODAY).subject == "Re: Application (Ref: A1)"
        assert assemble_letter(P, ["x"], LetterDetails(), TODAY).subject == ""

    def test_unnamed_recipient_gets_the_accepted_fallback_never_to_whom_it_may_concern(self):
        letter = assemble_letter(P, ["x"], LetterDetails(), TODAY)
        assert letter.salutation == "Dear Hiring Manager," and "whom it may concern" not in letter.text().lower()
        assert letter.recipient_lines == ()

    def test_availability_joins_the_last_paragraph_with_a_full_stop(self):
        letter = assemble_letter(P, ["One.", "Two."], LetterDetails(availability="I can start on 1 November"), TODAY)
        assert letter.paragraphs == ("One.", "Two. I can start on 1 November.")
        done = assemble_letter(P, ["One."], LetterDetails(availability="I can start now!"), TODAY)
        assert done.paragraphs == ("One. I can start now!",)

    def test_enclosure_can_be_left_out(self):
        assert "Enclosure" not in assemble_letter(P, ["x"], LetterDetails(enclosure=False), TODAY).text()

    def test_details_are_cleaned_and_a_hostile_recipient_name_is_refused(self):
        assert LetterDetails(company="  Acme   Corp \n").cleaned().company == "Acme Corp"
        assert len(LetterDetails(source="x" * 500).cleaned().source) == 80
        for bad in ("Ms. Rao <script>", "Rao\nBcc: evil@x.com", "1234", "{{x}}"):
            with pytest.raises(ValueError):
                LetterDetails(recipient_name=bad).cleaned()

    def test_sender_block_uses_only_data_the_candidate_has(self):
        bare = Profile(name="Sam Rivera")
        letter = assemble_letter(bare, ["x"], LetterDetails(), TODAY)
        assert letter.sender_lines == () and letter.text().startswith("Sam Rivera\n\n")


class TestEvidence:
    REPORT = analyze(P, JD, today=TODAY)

    def test_pairs_are_the_candidates_own_bullets_for_requirements_the_cv_really_meets(self):
        pairs = evidence_pairs(P, self.REPORT.matches)
        own = {b for e in P.experience for b in e.bullets}
        assert 1 <= len(pairs) <= 3 and all(e.bullet in own for e in pairs)
        assert all(e.requirement.lower() in e.bullet.lower() for e in pairs)
        assert len({e.bullet for e in pairs}) == len(pairs)  # a bullet is used once

    def test_a_skill_only_listed_under_skills_is_not_evidence(self):
        p = Profile(
            name="A B",
            skills=["Rust"],
            experience=[Experience(title="Dev", company="Co", bullets=["Built tools in Python for 40 users"])],
        )
        report = analyze(p, "Dev\nRequirements: Rust, Python", today=TODAY)
        assert [e.requirement for e in evidence_pairs(p, report.matches)] == ["Python"]

    def test_required_requirements_come_before_preferred_ones(self):
        jd = "Dev\nRequirements: Redis.\nNice to have: Kafka."
        pairs = evidence_pairs(P, analyze(P, jd, today=TODAY).matches)
        assert [e.requirement for e in pairs][:1] == ["Redis"]

    def test_measurable_action_bullets_win_between_candidates(self):
        p = Profile(
            name="A B",
            experience=[
                Experience(
                    title="Dev",
                    company="Acme",
                    bullets=["Worked with Kafka sometimes", "Cut Kafka lag by 40% using partitioning"],
                )
            ],
        )
        (pair,) = evidence_pairs(p, analyze(p, "Dev\nRequirements: Kafka", today=TODAY).matches)
        assert pair.bullet == "Cut Kafka lag by 40% using partitioning"
        assert pair.as_sentence() == "At Acme, I cut Kafka lag by 40% using partitioning."

    def test_project_evidence_is_labelled_as_a_project_not_an_employer(self):
        p = Profile(
            name="A B", projects=[Project(name="Tool", bullets=["Built a Kafka consumer handling 10k events/s"])]
        )
        (pair,) = evidence_pairs(p, analyze(p, "Dev\nRequirements: Kafka", today=TODAY).matches)
        assert (
            pair.employer == ""
            and pair.as_sentence() == "In my Tool project, I built a Kafka consumer handling 10k events/s."
        )

    def test_a_requirement_named_only_in_a_projects_tech_list_pairs_with_that_projects_bullet(self):
        p = Profile(
            name="A B",
            projects=[Project(name="Tracker", tech=["React"], bullets=["Built an expense tracker with charts"])],
        )
        (pair,) = evidence_pairs(p, analyze(p, "Dev\nRequirements: React", today=TODAY).matches)
        assert (pair.requirement, pair.project) == ("React", "Tracker")

    def test_a_skills_list_mention_with_no_supporting_bullet_gives_no_evidence(self):
        p = Profile(name="A B", skills=["Kafka"], projects=[Project(name="Tool", bullets=["Built a CLI"])])
        assert evidence_pairs(p, analyze(p, "Dev\nRequirements: Kafka", today=TODAY).matches) == []

    def test_no_matches_means_no_pairs_and_unconfirmed_items_are_ignored(self):
        assert evidence_pairs(P, ()) == []
        ghost = Experience(title="Wizard", company="FakeCo", origin="suggested", bullets=["Cut Kafka lag by 90%"])
        p = Profile(name="A B", experience=[ghost])
        assert evidence_pairs(p, analyze(p, "Dev\nRequirements: Kafka", today=TODAY).matches) == []


class TestLint:
    def good(self):
        return assemble_letter(P, GOOD_PARAS, replace(FULL, role="Senior Backend Engineer"), TODAY)

    def test_a_good_letter_passes_everything(self):
        results = lint_letter(self.good(), P, JD, company="Northwind Payments", role="Senior Backend Engineer", pages=1)
        assert [c.label for c in results if not c.passed] == []

    def test_every_check_can_fail_and_says_why(self):
        bad = assemble_letter(
            P,
            ["I am writing to apply for the role. I am a team player and a hard-working self-starter!!",
             "I love this. I like that. I want this. I enjoy that. I am passionate about it."],
            LetterDetails(), TODAY,
        )  # fmt: skip
        c = checks(bad, company="Northwind Payments", role="Senior Backend Engineer", pages=2)
        failed = {k for k, v in c.items() if not v.passed}
        assert {
            "length",
            "paragraphs",
            "one_page",
            "company",
            "role",
            "tailored",
            "evidence",
            "cliches",
            "variety",
            "cta",
        } <= failed
        assert (
            "'team player'" in c["cliches"].detail
            and "2 page" in c["one_page"].detail
            and "words" in c["length"].detail
        )

    @pytest.mark.parametrize(
        "words,ok,severity",
        [
            (100, False, "fix"),
            (200, False, "improve"),
            (250, True, "improve"),
            (400, True, "improve"),
            (430, False, "improve"),
            (500, False, "fix"),
        ],
    )
    def test_length_thresholds(self, words, ok, severity):
        first, rest = words - 2 * (words // 3), words // 3
        paras = [" ".join(["word"] * n) for n in (first, rest, rest)]
        letter = assemble_letter(P, paras, LetterDetails(), TODAY)
        assert len(letter.body.split()) == words
        c = lint_letter(letter, P, "")[0]
        assert (c.key, c.passed) == ("length", ok) and (ok or c.severity == severity)

    @pytest.mark.parametrize("phrase", CLICHES)
    def test_every_filler_phrase_is_caught_case_insensitively(self, phrase):
        letter = assemble_letter(P, ["Intro.", f"Some text. {phrase.upper()} here.", "Close."], LetterDetails(), TODAY)
        assert not checks(letter, job="")["cliches"].passed

    def test_variety_counts_sentences_that_start_with_i(self):
        many = assemble_letter(P, ["I did a. I did b. I did c. I did d.", "I did e.", "y."], LetterDetails(), TODAY)
        few = assemble_letter(P, ["Results came first. I did a. Teams grew.", "x.", "y."], LetterDetails(), TODAY)
        assert not checks(many, job="")["variety"].passed and checks(few, job="")["variety"].passed

    def test_a_closing_without_a_next_step_fails_cta(self):
        no = assemble_letter(P, ["a.", "b.", "Thanks."], LetterDetails(), TODAY)
        yes = assemble_letter(P, ["a.", "b.", "I would welcome a conversation."], LetterDetails(), TODAY)
        assert not checks(no, job="")["cta"].passed and checks(yes, job="")["cta"].passed

    def test_tone_flags_emoji_shouting_and_exclamation_marks(self):
        for text in ("Great news 🎉", "I am EXTREMELY keen", "Wow! Amazing! Yes!"):
            letter = assemble_letter(P, ["a.", text, "c."], LetterDetails(), TODAY)
            assert not checks(letter, job="")["tone"].passed, text

    def test_placeholders_fail_hard(self):
        letter = assemble_letter(P, ["a.", "At [Company] I did x.", "c."], LetterDetails(), TODAY)
        c = checks(letter, job="")["no_placeholders"]
        assert not c.passed and c.severity == "fix"

    @staticmethod
    def words_letter(n, paragraphs=3):
        """A letter whose body is exactly n words in the given number of paragraphs."""
        size, extra = divmod(n, paragraphs)
        sizes = [size + (1 if i < extra else 0) for i in range(paragraphs)]
        letter = assemble_letter(P, [" ".join(["word"] * k) for k in sizes], LetterDetails(), TODAY)
        assert len(letter.body.split()) == n
        return letter

    @pytest.mark.parametrize(
        "n,ok,severity",
        [(149, False, "fix"), (150, False, "improve"), (249, False, "improve"), (250, True, ""),
         (400, True, ""), (401, False, "improve"), (450, False, "improve"), (451, False, "fix")],
    )  # fmt: skip
    def test_word_count_boundaries_are_exact(self, n, ok, severity):
        c = lint_letter(self.words_letter(n), P, "")[0]
        assert c.passed is ok and (ok or c.severity == severity)

    @pytest.mark.parametrize("n,ok", [(2, False), (3, True), (5, True), (6, False)])
    def test_paragraph_count_boundaries_are_exact(self, n, ok):
        c = {x.key: x for x in lint_letter(self.words_letter(300, n), P, "")}["paragraphs"]
        assert c.passed is ok

    def test_i_start_ratio_of_exactly_half_passes_and_more_fails(self):
        half = assemble_letter(
            P, ["I did a. I did b. I did c.", "Results came. Teams grew. Costs fell."], LetterDetails(), TODAY
        )
        more = assemble_letter(P, ["I did a. I did b. I did c. I did d.", "Results came. Costs fell."], LetterDetails(), TODAY)  # fmt: skip
        assert "3 of 6" in checks(half, job="")["variety"].detail and checks(half, job="")["variety"].passed
        assert "4 of 6" in checks(more, job="")["variety"].detail and not checks(more, job="")["variety"].passed

    def test_tone_limits_are_one_exclamation_and_four_capital_letters(self):
        def tone(text):
            return checks(assemble_letter(P, ["a.", text, "c."], LetterDetails(), TODAY), job="")["tone"].passed

        assert tone("Delighted! Truly.") and not tone("Delighted! Truly! Yes.")
        assert tone("I used the NASA API.") and not tone("I used the STRIPE API.")

    def test_a_multi_word_role_needs_half_its_words(self):
        def role(text, title="Senior Backend Engineer"):
            letter = assemble_letter(P, ["a.", text, "c."], LetterDetails(), TODAY)
            return checks(letter, job="", role=title)["role"].passed

        assert role("As an engineer I build backend systems.") and not role("As a cook I bake bread.")
        assert role("I am a data scientist.", "Data Scientist") and not role("I am a chef.", "Data Scientist")
        four = "Senior Staff Backend Engineer"  # half of four words is two
        assert role("A senior engineer.", four) and not role("An engineer.", four)

    def test_tailoring_counts_required_skills_only_and_needs_two_when_two_exist(self):
        def tailored(text, job):
            letter = assemble_letter(P, ["a.", text, "c."], LetterDetails(), TODAY)
            return checks(letter, job=job)["tailored"].passed

        two = "Dev\nRequirements: Kafka, Redis"
        assert tailored("I used Kafka and Redis.", two) and not tailored("I used Kafka.", two)
        one = "Dev\nRequirements: Kafka"
        assert tailored("I used Kafka.", one) and not tailored("I used cooking.", one)
        three = "Dev\nRequirements: Kafka, Redis, Docker"
        assert tailored("I used Kafka and Redis.", three)  # two examples are enough even when three are asked for
        none = checks(assemble_letter(P, ["a.", "b.", "c."], LetterDetails(), TODAY), job=three)["tailored"]
        assert "0 of the 3" in none.detail and "none" in none.detail
        preferred = "Dev\nRequirements: Kafka\nNice to have: Redis, Docker, Terraform"
        assert tailored("I used Kafka.", preferred)  # preferred skills do not raise the bar

    def test_results_are_immutable(self):
        from dataclasses import FrozenInstanceError

        with pytest.raises(FrozenInstanceError):
            lint_letter(self.good(), P, "")[0].passed = False

    def test_every_check_can_both_pass_and_fail(self):
        # guards against "always true" entries: each key must pass on a good letter and fail on a bad one
        kw = {"company": "Northwind Payments", "role": "Senior Backend Engineer", "pages": 1}
        good = {c.key: c.passed for c in lint_letter(self.good(), P, JD, **kw)}
        bad = {
            c.key: c.passed
            for c in lint_letter(
                assemble_letter(P, ["I am writing to apply."], LetterDetails(), TODAY),
                P, JD, company="Zed", role="Quantum Chef", pages=2,
            )
        }  # fmt: skip
        assert set(good) == set(bad) and all(good.values())
        # tone and placeholders have their own dedicated failing tests above
        assert {k for k, ok in bad.items() if ok} <= {"tone", "no_placeholders"}


class TestDetailsAndHelpers:
    @pytest.mark.parametrize(
        "field,cap",
        [("company", 80), ("role", 100), ("recipient_title", 80), ("company_city", 60), ("job_ref", 40),
         ("source", 80), ("availability", 160), ("hook", 240)],
    )  # fmt: skip
    def test_every_free_text_field_is_capped_at_its_limit(self, field, cap):
        d = LetterDetails(**{field: "x" * (cap + 50)}).cleaned()
        assert len(getattr(d, field)) == cap
        assert len(getattr(LetterDetails(**{field: "y" * cap}).cleaned(), field)) == cap

    def test_recipient_name_cap_and_charset(self):
        assert len(LetterDetails(recipient_name="A" * 100).cleaned().recipient_name) == 60
        for bad in ("a@b", "x;y", "<b>", "Ms. Rao 42"):
            with pytest.raises(ValueError):
                LetterDetails(recipient_name=bad).cleaned()
        assert LetterDetails(recipient_name="Anne-Marie O'Neil Jr.").cleaned().recipient_name == "Anne-Marie O'Neil Jr."

    def test_details_are_immutable(self):
        from dataclasses import FrozenInstanceError

        with pytest.raises(FrozenInstanceError):
            LetterDetails().company = "x"

    def test_join_reads_like_prose(self):
        from careerforge.writing.basic_letter import _join

        assert (_join([]), _join(["a"]), _join(["a", "b"]), _join(["a", "b", "c"])) == (
            "",
            "a",
            "a and b",
            "a, b and c",
        )


class TestEvidenceRanking:
    @staticmethod
    def profile(*bullets, project_bullets=()):
        return Profile(
            name="A B",
            experience=[Experience(title="Dev", company="Acme", bullets=list(bullets))],
            projects=[Project(name="Side", bullets=list(project_bullets))] if project_bullets else [],
        )

    def pairs(self, p, job, limit=3):
        return evidence_pairs(p, analyze(p, job, today=TODAY).matches, limit=limit)

    def test_required_requirements_come_before_preferred_ones(self):
        p = self.profile("Built Redis caching for the API", "Built Kafka pipelines for events")
        job = "Dev\nRequirements: Kafka\nNice to have: Redis"
        assert [e.requirement for e in self.pairs(p, job)] == ["Kafka", "Redis"]

    def test_the_limit_is_honoured_and_each_bullet_is_used_once(self):
        p = self.profile("Built Kafka and Redis pipelines", "Built Docker images", "Ran Terraform stacks")
        job = "Dev\nRequirements: Kafka, Redis, Docker, Terraform"
        found = self.pairs(p, job, limit=2)
        assert len(found) == 2 and len({e.bullet for e in found}) == 2
        assert len(self.pairs(p, job, limit=3)) == 3

    def test_a_bullet_with_a_number_and_a_verb_beats_a_vague_one(self):
        p = self.profile("Worked with Kafka somewhat", "Cut Kafka lag by 40%")
        (e,) = self.pairs(p, "Dev\nRequirements: Kafka")
        assert e.bullet == "Cut Kafka lag by 40%"

    def test_a_short_bullet_beats_a_long_one_when_otherwise_equal(self):
        long = "Cut Kafka lag by 40% " + "across many regions " * 12
        p = self.profile(long, "Cut Kafka lag by 40%")
        (e,) = self.pairs(p, "Dev\nRequirements: Kafka")
        assert e.bullet == "Cut Kafka lag by 40%"

    def test_importance_orders_required_then_unspecified_then_preferred_not_by_quality(self):
        from careerforge.ats.jd import Requirement
        from careerforge.ats.match import Match

        p = self.profile("Built Alpha systems", "Built Beta systems", "Built Gamma systems")
        matches = tuple(
            Match(Requirement(term, "skill", importance), quality, ("experience",))
            for term, importance, quality in (("Alpha", "preferred", 1.0), ("Beta", "required", 0.1), ("Gamma", "unspecified", 0.5))
        )  # fmt: skip
        assert [e.requirement for e in evidence_pairs(p, matches)] == ["Beta", "Gamma", "Alpha"]
        assert [e.requirement for e in evidence_pairs(p, matches, limit=2)] == ["Beta", "Gamma"]

    def test_equal_importance_is_ordered_by_match_quality(self):
        from careerforge.ats.jd import Requirement
        from careerforge.ats.match import Match

        p = self.profile("Built Alpha systems", "Built Beta systems")
        matches = (
            Match(Requirement("Alpha", "skill", "required"), 0.4, ("experience",)),
            Match(Requirement("Beta", "skill", "required"), 0.9, ("experience",)),
        )
        assert [e.requirement for e in evidence_pairs(p, matches)] == ["Beta", "Alpha"]

    def test_three_pairs_is_the_default_limit(self):
        from careerforge.ats.jd import Requirement
        from careerforge.ats.match import Match

        names = ("Alpha", "Beta", "Gamma", "Delta")
        p = self.profile(*(f"Built {n}" for n in names))
        matches = tuple(Match(Requirement(t, "skill", "required"), 0.7, ("experience",)) for t in names)
        assert len(evidence_pairs(p, matches)) == 3

    def test_a_number_outweighs_an_action_verb_which_outweighs_brevity(self):
        number_only = "The Kafka team grew 3x"  # a number, no leading verb
        verb_only = "Designed Kafka consumers carefully"  # a leading verb, no number
        p = self.profile(verb_only, number_only)
        (e,) = self.pairs(p, "Dev\nRequirements: Kafka")
        assert e.bullet == number_only
        long_verb = ("Designed Kafka consumers " + "carefully and thoroughly " * 12).strip()  # >= 200 chars
        short_plain = "Kafka was used"
        assert len(long_verb) >= 200
        (e,) = self.pairs(self.profile(short_plain, long_verb), "Dev\nRequirements: Kafka")
        assert e.bullet == long_verb  # a verb (1.0) beats being short (0.5)

    def test_the_two_hundred_character_edge_for_brevity(self):
        def bullet(n):
            return "Kafka " + "x" * (n - 6)

        p = self.profile(bullet(200), bullet(199))
        (e,) = self.pairs(p, "Dev\nRequirements: Kafka")
        assert len(e.bullet) == 199

    def test_evidence_is_immutable(self):
        from dataclasses import FrozenInstanceError

        with pytest.raises(FrozenInstanceError):
            Evidence("a", "b", "c").bullet = "x"


class TestBasicLetter:
    def make(self, details=FULL, profile=P, job=JD):
        return build_basic_letter(profile, job, details=details, today=TODAY)

    def test_structure_and_content(self):
        draft = self.make()
        text = draft.text
        assert "Dear Ms. Rao," in text and "Re: Application for" in text and "(Ref: ENG-2041)" in text
        assert "I am applying for the Senior Backend Engineer" in draft.body and "through LinkedIn" in draft.body
        assert "Northwind's focus on reliable settlement" in draft.body  # the candidate's own hook, verbatim
        assert draft.verified and all(isinstance(e.bullet, str) for e in draft.evidence)

    def test_every_fact_is_the_candidates_own_so_the_claims_check_finds_nothing(self):
        draft = self.make()
        assert unsupported_claims(draft.body, profile_text=profile_text(P), years=8, posting_text=JD,
                                  extra_known="Northwind Payments LinkedIn Ms. Rao " + FULL.availability) == []  # fmt: skip

    def test_availability_comes_before_the_thanks(self):
        body = self.make().body
        assert body.index("I can start on 1 November.") < body.index("Thank you for your time")

    def test_without_a_hook_no_filler_paragraph_is_invented(self):
        body = self.make(replace(FULL, hook="")).body
        assert len(body.split("\n\n")) == 3  # opening, evidence, close
        assert not re.search(r"admire|passion|excited|love", body, re.I)

    def test_requirement_leads_vary_so_it_does_not_read_like_a_template(self):
        evidence = self.make().letter.paragraphs[1]
        assert "Your posting asks for" in evidence and "You also look for" in evidence

    def test_a_cv_with_no_matching_evidence_still_gives_a_short_honest_letter(self):
        draft = self.make(
            profile=Profile(name="Sam Rivera", contact=P.contact, skills=["Excel"]),
            job="Chef\nRequirements: knife skills",
        )
        assert draft.evidence == () and len(draft.letter.paragraphs) == 3
        assert draft.verified  # nothing unverifiable was claimed
        assert any(c.key == "length" and not c.passed for c in draft.checks)  # and the lint says it is thin

    def test_unconfirmed_ai_items_never_appear(self):
        ghost = Experience(title="Wizard", company="FakeCo", origin="suggested", bullets=["Cut Kafka lag by 90%"])
        draft = self.make(profile=P.model_copy(update={"experience": [ghost, *P.experience]}))
        assert "FakeCo" not in draft.text

    def test_opening_uses_the_headline_else_the_latest_title_and_omits_years_under_one(self):
        base = Profile(name="Sam Rivera", contact=P.contact, skills=["Python"])
        titled = base.model_copy(update={"experience": [Experience(title="Data Analyst", company="Acme")]})
        assert "I am a Data Analyst" in self.make(profile=titled).body
        assert "I bring experience" in self.make(profile=base).body
        assert "years of experience" not in self.make(profile=titled).body  # no dates: nothing to count
        assert "working mainly with Python" in self.make(profile=base).body

    def test_company_and_role_are_taken_from_the_posting_when_not_typed(self):
        draft = self.make(
            details=LetterDetails(), job="Senior Backend Engineer\nCompany: Zephyr Labs\nRequirements: Kafka"
        )
        assert "Zephyr Labs" in draft.body and "Senior Backend Engineer" in draft.body

    def test_without_a_company_or_role_the_wording_stays_grammatical(self):
        draft = self.make(details=LetterDetails(), job="Requirements: Kafka")
        assert "an opportunity" in draft.body and " at ." not in draft.body and "contribute." in draft.body

    def test_availability_gets_a_full_stop_only_when_it_lacks_one(self):
        no_stop = self.make(details=LetterDetails(company="Acme", availability="I can start soon")).body
        stop = self.make(details=LetterDetails(company="Acme", availability="I can start soon!")).body
        assert "I can start soon. Thank you" in no_stop and "I can start soon! Thank you" in stop

    def test_a_hook_gets_a_full_stop_only_when_it_lacks_one(self):
        assert "I like your mission." in self.make(details=LetterDetails(hook="I like your mission")).body

    def test_years_are_rounded_from_the_cv_dates_and_dropped_under_one(self):
        def worked(start, end):
            from careerforge.models import DateRange

            return Profile(
                name="Sam Rivera", contact=P.contact, headline="Analyst", skills=["Python"],
                experience=[Experience(title="Analyst", company="Acme", dates=DateRange(start=start, end=end))],
            )  # fmt: skip

        assert "with 2 years of experience" in self.make(profile=worked("2024-10", "2026-10")).body
        assert "years of experience" not in self.make(profile=worked("2026-06", "2026-10")).body
        assert "with 1 years of experience" in self.make(profile=worked("2025-10", "2026-10")).body

    def test_seven_and_a_half_years_of_dates_round_to_eight(self):
        from careerforge.models import DateRange

        profile = Profile(
            name="Sam Rivera", contact=P.contact, headline="Analyst", skills=["Python"],
            experience=[Experience(title="Analyst", company="Acme", dates=DateRange(start="2019-04", end="2026-10"))],
        )  # fmt: skip
        assert "with 8 years of experience" in self.make(profile=profile).body  # 91 months / 12 = 7.6

    def test_a_company_typed_by_the_user_wins_over_the_one_found_in_the_posting(self):
        draft = self.make(
            details=LetterDetails(company="Typed Co"), job="Backend Engineer\nCompany: Zephyr Labs\nRequirements: Kafka"
        )
        assert "Typed Co" in draft.body and "Zephyr" not in draft.body

    def test_at_most_three_skills_are_named_in_the_opening(self):
        profile = Profile(name="Sam Rivera", contact=P.contact, headline="Dev", skills=["A1", "B2", "C3", "D4", "E5"])
        body = self.make(profile=profile).body
        assert "working mainly with A1, B2 and C3." in body and "D4" not in body

    def test_default_date_is_today(self):
        assert format_date(date.today(), "us") in build_basic_letter(P, JD).letter.date_text

    def test_passes_the_structural_lint_for_a_full_cv(self):
        c = {x.key: x for x in self.make().checks}
        assert all(
            c[k].passed
            for k in ("paragraphs", "company", "role", "tailored", "evidence", "cliches", "cta", "no_placeholders")
        )

    def test_fresher_with_projects_gets_project_evidence(self):
        draft = self.make(
            profile=fresher(), job="Dev\nRequirements: React, MongoDB", details=LetterDetails(company="Acme")
        )
        assert "In my Expense Tracker project, I built a full-stack web app" in draft.body
        assert [e.project for e in draft.evidence] == ["Expense Tracker"]


class TestAiLetter:
    def write(self, *replies, **kw):
        fake = Scripted(*replies)
        return write_cover_letter(client(fake), P, JD, today=TODAY, **kw), fake

    def test_prompt_carries_evidence_tone_hook_source_and_bans_filler(self):
        draft, fake = self.write({"paragraphs": GOOD_PARAS}, details=replace(FULL, tone="concise"))
        req = fake.requests[0]
        assert "Evidence to use" in req.user and 'Posting asks for "' in req.user
        assert "Reason for this company" in req.user and "Northwind's focus on reliable settlement" in req.user
        assert "found the role via: LinkedIn" in req.user and "brisk and direct" in req.system
        assert '"team player"' in req.system and "250 to 350 words" in req.system
        assert draft.evidence and draft.letter.salutation == "Dear Ms. Rao,"

    def test_letter_is_linted_and_a_good_one_passes_the_important_checks(self):
        draft, _ = self.write({"paragraphs": GOOD_PARAS}, details=FULL)
        assert draft.verified
        assert {c.key for c in draft.checks if not c.passed} <= {"variety"}  # at worst a style tip

    def test_source_recipient_and_hook_are_not_flagged_as_invented_names(self):
        paras = list(GOOD_PARAS)
        paras[0] = "I found your Senior Backend Engineer role on LinkedIn. " + paras[0]
        draft, _ = self.write({"paragraphs": paras}, details=FULL)
        assert draft.verified

    def test_a_non_person_recipient_name_is_refused_before_any_ai_call(self):
        fake = Scripted()
        with pytest.raises(ValueError):
            write_cover_letter(client(fake), P, JD, today=TODAY, details=LetterDetails(recipient_name="<b>x</b>"))
        assert fake.requests == []

    def test_old_keyword_arguments_still_work(self):
        draft, _ = self.write({"paragraphs": GOOD_PARAS}, addressee="Ms. Rao", company="Northwind Payments")
        assert "Dear Ms. Rao," in draft.text

    def test_too_few_or_too_many_paragraphs_are_rejected_and_repaired(self):
        draft, fake = self.write({"paragraphs": GOOD_PARAS[:2]}, {"paragraphs": GOOD_PARAS})
        assert draft.verified and len(fake.requests) == 2 and len(draft.letter.paragraphs) == len(GOOD_PARAS)

    def test_system_prompt_formats(self):
        rendered = SYSTEM.format(cliches='"a", "b"')
        assert '"a", "b"' in rendered and '{"paragraphs"' in rendered and "{cliches}" not in rendered


class TestRendering:
    LETTER = build_basic_letter(P, JD, details=FULL, today=TODAY).letter

    def test_pdf_is_one_page_with_selectable_text_and_correct_metadata(self):
        for template in ("ats-classic", "ats-modern", "executive", "sidebar"):
            out = render_letter_pdf(self.LETTER, RenderOptions(template=template))
            reader = PdfReader(io.BytesIO(out.data))
            text = " ".join(page.extract_text() for page in reader.pages)
            assert out.pages == 1 and out.warnings == ()
            assert (
                "Dear Ms. Rao," in text
                and "Priya Nair" in text
                and "Yours sincerely" not in text.replace("Sincerely", "")
            )
            assert reader.metadata.title == "Priya Nair - Cover letter" and reader.metadata.author == "Priya Nair"

    def test_a_long_letter_is_flagged_as_two_pages(self):
        long = replace(self.LETTER, paragraphs=("A sentence about delivery and results. " * 60,) * 6)
        out = render_letter_pdf(long)
        assert out.pages > 1 and "second page" in out.warnings[0]

    def test_paper_and_density_options_apply(self):
        a4 = PdfReader(io.BytesIO(render_letter_pdf(self.LETTER, RenderOptions(paper="A4")).data)).pages[0].mediabox
        letter = (
            PdfReader(io.BytesIO(render_letter_pdf(self.LETTER, RenderOptions(paper="Letter")).data)).pages[0].mediabox
        )
        assert round(float(a4.height)) == 842 and round(float(letter.height)) == 792

    def test_html_escapes_everything_and_forbids_scripts(self):
        evil = replace(self.LETTER, paragraphs=('5 < 6 && "x" <img src=x onerror=alert(1',), subject="a </style b")
        html = render_letter_html(evil)
        assert "<img" not in html and "&lt;img" in html and "default-src &#39;none&#39;" in html
        assert html.count("<style>") == 1 and "</style b" not in html

    def test_docx_uses_real_paragraphs_in_order_with_no_tables(self):
        d = docx.Document(io.BytesIO(render_letter_docx(self.LETTER, RenderOptions(template="ats-modern"))))
        texts = [p.text for p in d.paragraphs]
        assert not d.tables and texts[0] == "Priya Nair" and "Dear Ms. Rao," in texts and "Sincerely," in texts
        assert texts.index("Dear Ms. Rao,") < texts.index("Sincerely,")
        assert d.core_properties.author == "Priya Nair" and "python-docx" not in d.core_properties.author

    def test_txt_is_the_letter_text(self):
        assert render_letter_txt(self.LETTER) == self.LETTER.text()

    @pytest.mark.parametrize(
        "mutate,fragment",
        [
            (lambda x: replace(x, paragraphs=("At [Company] I did x.",)), "placeholder"),
            (lambda x: replace(x, paragraphs=()), "no body"),
            (lambda x: replace(x, sender_lines=()), "contact"),
            (lambda x: replace(x, sender_name=""), "name"),
            (lambda x: replace(x, paragraphs=("Hello {{name}}",)), "placeholder"),
        ],
    )
    def test_unfinished_letters_are_refused_in_every_format(self, mutate, fragment):
        bad = mutate(self.LETTER)
        for fn in (render_letter_html, render_letter_pdf, render_letter_docx, render_letter_txt):
            with pytest.raises(LetterNotReady, match=fragment):
                fn(bad)

    def test_edited_body_replaces_paragraphs_on_blank_lines(self):
        edited = with_body(self.LETTER, "First para\nstill first.\n\n\nSecond para.\n")
        assert (
            edited.paragraphs == ("First para still first.", "Second para.")
            and edited.salutation == self.LETTER.salutation
        )
