"""DOCX -> text. Includes tables, header text (contact details often live there), text boxes and hyperlinks.

Legacy .doc and macro-enabled .docm are refused. Archive size/entry limits guard against zip bombs before
python-docx parses anything.
"""

from __future__ import annotations

import io
import zipfile

from docx import Document as load_docx
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from careerforge.ingest.errors import UnreadableDocument
from careerforge.text import clean_text

MAX_UNCOMPRESSED = 60_000_000
MAX_ENTRIES = 3000


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > MAX_ENTRIES or sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
                raise UnreadableDocument("This Word file is unusually large inside; it was not opened for safety")
            names = {i.filename for i in infos}
            if "word/document.xml" not in names:
                raise UnreadableDocument("This does not look like a Word (.docx) document")
            if "word/vbaProject.bin" in names:
                raise UnreadableDocument("Macro-enabled Word files are not accepted; save a plain .docx instead")
    except zipfile.BadZipFile as exc:
        raise UnreadableDocument("This Word file could not be opened; it may be damaged") from exc


def _table_lines(table: Table) -> list[str]:
    lines = []
    for row in table.rows:
        seen: list[str] = []
        for cell in row.cells:
            text = clean_text("\n".join(p.text for p in cell.paragraphs), multiline=True)
            if text and text not in seen:  # merged cells repeat the same text
                seen.append(text)
        if seen:
            lines.append(" | ".join(t.replace("\n", " ") for t in seen))
    return lines


def extract_docx(data: bytes) -> tuple[str, list[str]]:
    """Return (text, hyperlinks)."""
    _check_archive(data)
    try:
        doc = load_docx(io.BytesIO(data))
    except Exception as exc:  # python-docx raises assorted errors on malformed XML
        raise UnreadableDocument("This Word file could not be read; it may be damaged") from exc

    parts: list[str] = []
    links: list[str] = []
    for section in doc.sections[:1]:  # header of the first section: names/emails are often here
        parts += [p.text for p in section.header.paragraphs if p.text.strip()]

    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            para = Paragraph(child, doc)
            parts.append(para.text)
            links += [h.address for h in para.hyperlinks if h.address]
        elif child.tag == qn("w:tbl"):
            parts += _table_lines(Table(child, doc))

    # text boxes are invisible to python-docx's paragraph API
    in_body = "\n".join(parts)
    for box in body.iter(qn("w:txbxContent")):
        text = clean_text("\n".join("".join(t.text or "" for t in p.iter(qn("w:t"))) for p in box.iter(qn("w:p"))),
                          multiline=True)  # fmt: skip
        if text and text not in in_body:
            parts.append(text)
            in_body += "\n" + text

    uniq = [u for i, u in enumerate(links) if u.lower().startswith(("http://", "https://")) and u not in links[:i]]
    return clean_text("\n".join(parts), multiline=True), uniq[:20]
