"""Render a cover letter as HTML, PDF, DOCX or text, in the same look as the CV (template, accent, font, paper).

The PDF is built with the same rules as the CV: no network access, and the page count is reported because a cover
letter that runs onto a second page is the commonest layout complaint.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, replace

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from careerforge.render.cv import CSP, _env, _pdf_engine
from careerforge.render.docx_export import DOCX_FONT, PAGE_MM, _rgb
from careerforge.render.options import FONT_STACKS, RenderOptions
from careerforge.writing.letter_format import Letter

_PLACEHOLDER = re.compile(r"\[[^\]]*\]|\{[^}]*\}|<[A-Za-z/][^>]*>|\bTODO\b|\bXXX\b|lorem ipsum", re.I)
LETTER_SIZE = {"compact": (10.5, 1.35), "normal": (11, 1.45), "relaxed": (11.5, 1.55)}
LETTER_MARGIN_MM = {"compact": 20, "normal": 22, "relaxed": 25}


class LetterNotReady(ValueError):
    """The letter still has placeholders, no body or no sender."""


def require_ready(letter: Letter) -> Letter:
    problems = []
    if not letter.sender_name.strip():
        problems.append("the sender's name is missing")
    if not letter.sender_lines:
        problems.append("there are no contact details")
    if not letter.paragraphs or not any(p.strip() for p in letter.paragraphs):
        problems.append("the letter has no body")
    if _PLACEHOLDER.search(letter.text()):
        problems.append("it still contains placeholder text in brackets or markup")
    if problems:
        raise LetterNotReady("The letter is not ready: " + "; ".join(problems) + ".")
    return letter


def with_body(letter: Letter, text: str) -> Letter:
    """The letter with its body replaced by edited text (paragraphs are separated by blank lines)."""
    paragraphs = tuple(" ".join(block.split()) for block in re.split(r"\n\s*\n", text.strip()) if block.strip())
    return replace(letter, paragraphs=paragraphs)


def render_letter_html(letter: Letter, options: RenderOptions | None = None) -> str:
    options = options or RenderOptions()
    size, line_height = LETTER_SIZE[options.density]
    return (
        _env()
        .get_template("letter.html.j2")
        .render(
            letter=require_ready(letter),
            template=options.template if options.template != "sidebar" else "ats-classic",
            lang=options.lang,
            accent=options.accent,
            font_stack=FONT_STACKS[options.font_key],
            fs=size,
            lh=line_height,
            margin=LETTER_MARGIN_MM[options.density],
            paper=options.paper,
            csp=CSP,
        )
    )


@dataclass(frozen=True)
class LetterPdf:
    data: bytes
    pages: int
    warnings: tuple[str, ...] = ()


def render_letter_pdf(letter: Letter, options: RenderOptions | None = None) -> LetterPdf:
    html_class, no_fetch = _pdf_engine()
    html = render_letter_html(letter, options)
    document = html_class(string=html, url_fetcher=no_fetch).render()
    pages = len(document.pages)
    warnings = ("The letter runs onto a second page. Shorten it to one page.",) if pages > 1 else ()
    return LetterPdf(document.write_pdf(), pages, warnings)


def render_letter_txt(letter: Letter) -> str:
    return require_ready(letter).text()


def render_letter_docx(letter: Letter, options: RenderOptions | None = None) -> bytes:
    options = options or RenderOptions()
    require_ready(letter)
    size, line_height = LETTER_SIZE[options.density]
    margin = LETTER_MARGIN_MM[options.density]
    font = DOCX_FONT[options.font_key]
    doc = Document()
    section = doc.sections[0]
    width, height = PAGE_MM[options.paper]
    section.page_width, section.page_height = Mm(width), Mm(height)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = Mm(margin)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(size)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = line_height
    accent = _rgb(options.accent) if options.template == "ats-modern" else RGBColor(0, 0, 0)
    centred = options.template == "executive"

    def para(text: str = "", *, bold: bool = False, before: float = 0, after: float = 0, color=None, scale=1.0):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        run.font.size = Pt(size * scale)
        if color is not None:
            run.font.color.rgb = color
        p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
        if centred and before == 0 and after == 0 and scale != 1.0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    head = para(letter.sender_name, bold=True, color=accent, scale=1.6)
    if centred:
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in letter.sender_lines:
        line_p = para(line, scale=0.95)
        if centred:
            line_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para(letter.date_text, before=16)
    for k, line in enumerate(letter.recipient_lines):
        para(line, before=10 if k == 0 else 0)
    if letter.subject:
        para(letter.subject, bold=True, before=10)
    para(letter.salutation, before=10)
    for body in letter.paragraphs:
        para(body, before=8)
    para(letter.closing, before=12)
    para(letter.sender_name, bold=True, before=22)
    if letter.enclosure:
        para(letter.enclosure, before=10, scale=0.95)
    props = doc.core_properties
    props.title, props.author, props.last_modified_by = (
        f"{letter.sender_name} - Cover letter",
        letter.sender_name,
        letter.sender_name,
    )
    props.comments = props.subject = props.keywords = ""
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
